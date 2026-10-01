#!/usr/bin/env python3
"""Tests for cv-tailor improvements."""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

import pytest

# Add scripts to path
CV_TAILOR_DIR = Path(__file__).resolve().parents[2] / 'cv-tailor'
sys.path.insert(0, str(CV_TAILOR_DIR / 'scripts'))


# ── Task 2: build_analysis profile loading ──────────────────────────────────

class TestUserProfileLoading:
    def test_missing_profile_returns_empty_dict(self, monkeypatch, tmp_path):
        """Missing user-profile.yaml → graceful empty dict, no crash."""
        import build_analysis
        monkeypatch.setattr(build_analysis, 'PROFILE_PATH', tmp_path / 'nonexistent.yaml')
        result = build_analysis._load_user_profile()
        assert result == {}

    def test_invalid_yaml_returns_empty_dict(self, monkeypatch, tmp_path):
        """Invalid YAML → empty dict, no crash."""
        import build_analysis
        bad = tmp_path / 'user-profile.yaml'
        bad.write_text(': invalid: yaml: [[[', encoding='utf-8')
        monkeypatch.setattr(build_analysis, 'PROFILE_PATH', bad)
        result = build_analysis._load_user_profile()
        assert result == {}

    def test_valid_profile_is_injected(self, monkeypatch, tmp_path):
        """Valid profile → user_profile key present in pending-analysis.json."""
        import build_analysis

        profile_data = {
            'identity': {'summary_opener': 'Enterprise auditor', 'credential_label': 'FMVA charterholder', 'lead_identity': 'AI builder'},
            'domains': {'confirmed_asset_classes': ['hedge funds', 'real estate funds'], 'role_verb_map': {'Acme Audit': 'audited'}},
            'confidentiality': {'redact_clients': True, 'approved_names': []},
            'projects': {'poc_disclaimer': True, 'project_role_map': {}},
        }
        profile_file = tmp_path / 'user-profile.yaml'
        import yaml
        profile_file.write_text(yaml.dump(profile_data), encoding='utf-8')

        monkeypatch.setattr(build_analysis, 'PROFILE_PATH', profile_file)
        monkeypatch.setattr(build_analysis, 'DATA_DIR', tmp_path)
        monkeypatch.setattr(build_analysis, 'PENDING_PATH', tmp_path / 'pending-analysis.json')
        monkeypatch.setattr(build_analysis, 'ANALYSIS_PATH', tmp_path / 'analysis.json')
        monkeypatch.setattr(build_analysis, 'TERM_MAP_PATH', tmp_path / 'nonexistent.md')

        # Minimal fake docx
        from docx import Document
        fake_docx = tmp_path / 'resume.docx'
        doc = Document()
        doc.add_paragraph('Professional summary line.')
        doc.save(str(fake_docx))

        build_analysis.build('Acme', 'AI Architect', str(fake_docx), 'Job description text.')

        payload = json.loads((tmp_path / 'pending-analysis.json').read_text())
        assert 'user_profile' in payload
        assert payload['user_profile']['identity']['summary_opener'] == 'Enterprise auditor'
        assert payload['user_profile']['confidentiality']['redact_clients'] is True


# ── Task 7: docx_safe_patch fixes ───────────────────────────────────────────

class TestDocxSafePatchFixes:
    def _make_bold_header_doc(self, tmp_path) -> 'Path':
        """Create a docx with a bold-header paragraph: 'Header: continuation text'"""
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        p = doc.add_paragraph()
        run_bold = p.add_run('Reduced close cycle: ')
        run_bold.bold = True
        run_bold.font.name = 'Calibri'
        run_bold.font.size = Pt(11)
        run_normal = p.add_run('by 30% through process improvements.')
        run_normal.bold = False
        run_normal.font.name = 'Calibri'
        run_normal.font.size = Pt(11)
        path = tmp_path / 'test_bold.docx'
        doc.save(str(path))
        return path

    def test_bold_does_not_bleed_into_continuation(self, tmp_path):
        """Replacing a bold-header paragraph must not make entire line bold."""
        import docx_safe_patch
        from docx import Document

        src = self._make_bold_header_doc(tmp_path)
        dst = tmp_path / 'out_bold.docx'

        repls = [{'old': 'Reduced close cycle: by 30% through process improvements.',
                  'new': 'Accelerated close cycle: cut period-end processing time by 35%.'}]
        docx_safe_patch.apply_safe_patch(src, dst, repls)

        doc = Document(str(dst))
        para = next(p for p in doc.paragraphs if 'Accelerated close cycle' in p.text)
        non_bold_runs = [r for r in para.runs if not r.bold and r.text.strip()]
        assert len(non_bold_runs) > 0, "All runs are bold — bold bleed detected"

    def test_replacement_preserves_font_size(self, tmp_path):
        """Replacement run must inherit font size from original."""
        import docx_safe_patch
        from docx import Document

        src = self._make_bold_header_doc(tmp_path)
        dst = tmp_path / 'out_size.docx'

        repls = [{'old': 'by 30% through process improvements.',
                  'new': 'by 35% through standardised workflows.'}]
        docx_safe_patch.apply_safe_patch(src, dst, repls)

        doc = Document(str(dst))
        para = next(p for p in doc.paragraphs if '35%' in p.text)
        runs_with_size = [r for r in para.runs if r.font.size is not None]
        assert len(runs_with_size) > 0, "No runs have explicit font size set"


# ── Task 8: quality_gate additions ──────────────────────────────────────────

class TestQualityGateAdditions:
    def _make_doc(self, tmp_path, extra_text: str) -> 'Path':
        from docx import Document
        doc = Document()
        for section in ['PROFESSIONAL SUMMARY', 'TECHNICAL CAPABILITIES',
                        'TECHNICAL PROJECTS', 'PROFESSIONAL EXPERIENCE', 'EDUCATION']:
            doc.add_paragraph(section)
        doc.add_paragraph(extra_text)
        for i in range(20):
            doc.add_paragraph(f'Bullet point {i}.')
        path = tmp_path / 'resume.docx'
        doc.save(str(path))
        return path

    def _make_redline(self, tmp_path) -> 'Path':
        from docx import Document
        doc = Document()
        doc.add_paragraph('Redline.')
        path = tmp_path / 'redline.docx'
        doc.save(str(path))
        return path

    def test_buzzword_detected_as_failure(self, tmp_path):
        import quality_gate
        resume = self._make_doc(tmp_path, 'Leveraging AI tools to spearhead innovation.')
        redline = self._make_redline(tmp_path)
        result = quality_gate.qc(str(resume), str(redline))
        assert result['status'] == 'fail'
        assert len(result['buzzwords_found']) > 0

    def test_em_dash_detected_as_failure(self, tmp_path):
        import quality_gate
        resume = self._make_doc(tmp_path, 'Led finance ops — delivered results.')
        redline = self._make_redline(tmp_path)
        result = quality_gate.qc(str(resume), str(redline))
        assert result['status'] == 'fail'
        assert result['em_dash_found'] is True

    def test_clean_doc_passes(self, tmp_path):
        import quality_gate
        resume = self._make_doc(tmp_path, 'Led finance operations and delivered results.')
        redline = self._make_redline(tmp_path)
        result = quality_gate.qc(str(resume), str(redline))
        assert result['status'] == 'pass'


# ── Task 9: validate_analysis additions ─────────────────────────────────────

class TestValidateAnalysisAdditions:
    def _make_docx(self, tmp_path):
        from docx import Document
        p = tmp_path / 'resume.docx'
        Document().save(str(p))
        return str(p)

    def _base(self, tmp_path):
        return {
            'company': 'Acme',
            'role': 'AI Architect',
            'base_resume_path': self._make_docx(tmp_path),
            'summary_edits': [{'old': 'x', 'new': 'y'}],
            'bullet_edits': [],
            'keyword_targets': ['kw'],
            'cover_letter_paragraphs': ['p1', 'p2', 'p3'],
            'claims_guardrail': ['Do not invent.'],
        }

    def test_task_phrase_strength_emits_warning_not_error(self, tmp_path):
        import validate_analysis
        obj = self._base(tmp_path)
        obj['core_strengths'] = ['Managing investment funds', 'LLM Workflow Automation']
        errs, warns = validate_analysis.validate(obj)
        assert errs == [], f"Should be no errors, got: {errs}"
        assert any('strengths' in w.lower() or 'task' in w.lower() for w in warns)

    def test_client_name_rejected_when_redact_true(self, tmp_path):
        import validate_analysis
        obj = self._base(tmp_path)
        obj['bullet_edits'] = [{'old': 'x', 'new': 'Audited Morgan Stanley investment portfolio.'}]
        obj['confidentiality'] = {'redact_clients': True, 'approved_names': []}
        errs, warns = validate_analysis.validate(obj)
        assert any('morgan stanley' in e.lower() for e in errs)

    def test_approved_client_name_passes(self, tmp_path):
        import validate_analysis
        obj = self._base(tmp_path)
        obj['bullet_edits'] = [{'old': 'x', 'new': 'Audited Morgan Stanley investment portfolio.'}]
        obj['confidentiality'] = {'redact_clients': True, 'approved_names': ['Morgan Stanley']}
        errs, warns = validate_analysis.validate(obj)
        assert not any('morgan stanley' in e.lower() for e in errs)

    def test_existing_validation_still_returns_tuple(self, tmp_path):
        import validate_analysis
        obj = {'company': '', 'role': 'x', 'base_resume_path': self._make_docx(tmp_path),
               'summary_edits': [], 'bullet_edits': [], 'keyword_targets': [],
               'cover_letter_paragraphs': ['p1', 'p2', 'p3'], 'claims_guardrail': ['x']}
        result = validate_analysis.validate(obj)
        assert isinstance(result, tuple) and len(result) == 2
        errs, warns = result
        assert any('empty' in e for e in errs)
