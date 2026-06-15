#!/usr/bin/env python3
"""Tests for the cv-tailor honesty gates.

Covers:
- claims_gate.py: invented numbers/employers/credentials/dates caught,
  legitimate restatements pass, cover letter is scanned.
- generate_redline.py: redline includes all four applied edit categories.
- quality_gate.qc_cover_letter: cover letter gate failure paths, including
  the body-less letter boundary (scaffold alone must fail).
- validate_analysis: empty/whitespace cover letter paragraphs rejected.
- run_pipeline.cmd_apply: analysis.json survives failed runs.
"""
from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import pytest

CV_TAILOR_DIR = Path(__file__).resolve().parents[2] / 'cv-tailor'
sys.path.insert(0, str(CV_TAILOR_DIR / 'scripts'))


def _load_cv_run_pipeline():
    """Load cv-tailor's run_pipeline under a unique module name.

    job-search also has a run_pipeline.py; whichever test imports first wins
    the plain `run_pipeline` slot in sys.modules, so we avoid it entirely.
    """
    name = 'cv_tailor_run_pipeline'
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, CV_TAILOR_DIR / 'scripts' / 'run_pipeline.py')
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    saved_path = list(sys.path)
    try:
        spec.loader.exec_module(mod)
    finally:
        sys.path[:] = saved_path
    return mod

BASE_CV_TEXT = '\n'.join([
    'PROFESSIONAL SUMMARY',
    'Fund accountant with 7 years of experience.',
    'Managed $50M AUM across 12 investment funds.',
    'Cut close cycle time by 30% at Apex Group since 2019.',
])


def _analysis(**overrides) -> dict:
    obj = {
        'company': 'Acme Fund Services',
        'role': 'Senior Fund Accountant',
        'summary_edits': [],
        'bullet_edits': [],
        'tailored_edits': [],
        'shared_edits': [],
        'cover_letter_paragraphs': [],
    }
    obj.update(overrides)
    return obj


# ── Claims gate: deterministic no-invention checks ───────────────────────────

class TestClaimsGate:
    def test_invented_number_caught(self):
        import claims_gate
        obj = _analysis(bullet_edits=[{'old': 'Managed $50M AUM across 12 investment funds.',
                                       'new': 'Managed $500M AUM across 12 investment funds.'}])
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        assert result['status'] == 'fail'
        msgs = [v['message'] for v in result['violations']]
        assert any('$500M' in m and 'not found in base CV or profile' in m for m in msgs)

    def test_invented_employer_caught(self):
        import claims_gate
        obj = _analysis(tailored_edits=[{'old': 'x', 'new': 'Led NAV reviews for Goldman Sachs portfolios.'}])
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        assert result['status'] == 'fail'
        assert any(v['type'] == 'employer' and v['claim'] == 'goldman sachs'
                   for v in result['violations'])

    def test_invented_credential_caught(self):
        import claims_gate
        obj = _analysis(summary_edits=[{'old': 'x', 'new': 'CFA charterholder and fund accountant.'}])
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        assert result['status'] == 'fail'
        assert any(v['type'] == 'credential' and v['claim'] == 'CFA' for v in result['violations'])

    def test_invented_date_caught(self):
        import claims_gate
        obj = _analysis(bullet_edits=[{'old': 'x', 'new': 'Leading fund operations since 2015.'}])
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        assert result['status'] == 'fail'
        assert any(v['type'] == 'date' and '2015' in v['claim'] for v in result['violations'])

    def test_legitimate_restatement_passes(self):
        import claims_gate
        obj = _analysis(
            summary_edits=[{'old': 'x', 'new': 'Oversaw $50 million in AUM and reduced close time by 30 percent.'}],
            cover_letter_paragraphs=['In 7 years across 12 funds I built repeatable close processes.'],
        )
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        assert result['status'] == 'pass', f"Unexpected violations: {result['violations']}"

    def test_cover_letter_is_scanned(self):
        import claims_gate
        obj = _analysis(cover_letter_paragraphs=[
            'I am excited to apply.',
            'I have managed $2.5B in fund assets.',
            'Thank you for your consideration.',
        ])
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        assert result['status'] == 'fail'
        assert any('cover_letter_paragraphs[1]' in v['category'] and '$2.5B' in v['claim']
                   for v in result['violations'])

    def test_all_four_edit_categories_scanned(self):
        import claims_gate
        obj = _analysis(
            summary_edits=[{'old': 'x', 'new': 'Invented 11% figure.'}],
            bullet_edits=[{'old': 'x', 'new': 'Invented 22% figure.'}],
            tailored_edits=[{'old': 'x', 'new': 'Invented 33% figure.'}],
            shared_edits=[{'old': 'x', 'new': 'Invented 44% figure.'}],
        )
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        locations = {v['category'].split('[')[0] for v in result['violations']}
        assert locations == {'summary_edits', 'bullet_edits', 'tailored_edits', 'shared_edits'}

    def test_profile_sourced_claim_passes(self):
        import claims_gate
        profile = {'identity': {'credential_label': 'FMVA charterholder'}}
        obj = _analysis(summary_edits=[{'old': 'x', 'new': 'FMVA charterholder turned AI builder.'}])
        assert claims_gate.check_claims(obj, BASE_CV_TEXT, profile)['status'] == 'pass'
        assert claims_gate.check_claims(obj, BASE_CV_TEXT, {})['status'] == 'fail'

    def test_target_company_name_allowed(self):
        import claims_gate
        obj = _analysis(company='Vistra',
                        cover_letter_paragraphs=['I would love to join Vistra and contribute.'])
        result = claims_gate.check_claims(obj, BASE_CV_TEXT)
        assert result['status'] == 'pass', f"Unexpected violations: {result['violations']}"


# ── Redline completeness: all four applied categories ────────────────────────

class TestRedlineAllCategories:
    def test_redline_contains_all_four_categories(self, tmp_path):
        import generate_redline
        from docx import Document

        analysis = _analysis(
            summary_edits=[{'old': 'old summary', 'new': 'NEW-SUMMARY-TEXT'}],
            bullet_edits=[{'old': 'old bullet', 'new': 'NEW-BULLET-TEXT'}],
            tailored_edits=[{'old': 'old tailored', 'new': 'NEW-TAILORED-TEXT'}],
            shared_edits=[{'old': 'old shared', 'new': 'NEW-SHARED-TEXT'}],
        )
        analysis_path = tmp_path / 'analysis.json'
        analysis_path.write_text(json.dumps(analysis), encoding='utf-8')
        out = tmp_path / 'redline.docx'

        generate_redline.generate('base.docx', str(analysis_path), str(out))

        text = '\n'.join(p.text for p in Document(str(out)).paragraphs)
        for marker in ('NEW-SUMMARY-TEXT', 'NEW-BULLET-TEXT', 'NEW-TAILORED-TEXT', 'NEW-SHARED-TEXT'):
            assert marker in text, f'redline missing edit: {marker}'
        for cat in ('summary_edits', 'bullet_edits', 'tailored_edits', 'shared_edits'):
            assert cat in text, f'redline missing category label: {cat}'
        assert 'Total planned edits: 4' in text


# ── Cover letter quality gate ─────────────────────────────────────────────────

class TestCoverLetterGate:
    def _build_letter(self, tmp_path, paragraphs):
        run_pipeline = _load_cv_run_pipeline()
        path = tmp_path / 'cover.docx'
        run_pipeline._build_cover_letter(path, 'Acme', 'Senior Fund Accountant', paragraphs)
        return path

    def test_buzzword_fails_cover_letter_gate(self, tmp_path):
        import quality_gate
        path = self._build_letter(tmp_path, [
            'I am leveraging my background in fund accounting.',
            'Second paragraph.',
            'Third paragraph.',
        ])
        result = quality_gate.qc_cover_letter(str(path))
        assert result['status'] == 'fail'
        assert 'leveraging' in result['buzzwords_found']

    def test_em_dash_fails_cover_letter_gate(self, tmp_path):
        import quality_gate
        path = self._build_letter(tmp_path, [
            'Fund accounting — my core strength.',
            'Second paragraph.',
            'Third paragraph.',
        ])
        result = quality_gate.qc_cover_letter(str(path))
        assert result['status'] == 'fail'
        assert result['em_dash_found'] is True

    def test_clean_cover_letter_passes(self, tmp_path):
        import quality_gate
        path = self._build_letter(tmp_path, [
            'I am writing to apply for the role.',
            'My background fits the team well.',
            'Thank you for your consideration.',
        ])
        result = quality_gate.qc_cover_letter(str(path))
        assert result['status'] == 'pass', f"Unexpected reasons: {result['reasons']}"

    def test_empty_body_fails_cover_letter_gate(self, tmp_path):
        """The scaffold alone is 6 non-empty lines; a letter with zero body
        paragraphs must fail the minimum-content check."""
        import quality_gate
        path = self._build_letter(tmp_path, ['', '', ''])
        result = quality_gate.qc_cover_letter(str(path))
        assert result['status'] == 'fail'
        assert any('too short' in r for r in result['reasons'])

    def test_single_body_paragraph_passes_minimum(self, tmp_path):
        import quality_gate
        path = self._build_letter(tmp_path, ['One real body paragraph about the role.'])
        result = quality_gate.qc_cover_letter(str(path))
        assert result['status'] == 'pass', f"Unexpected reasons: {result['reasons']}"


# ── Analysis validation: cover letter paragraphs must have content ───────────

class TestValidateAnalysisCoverLetter:
    def test_empty_paragraphs_rejected(self, tmp_path):
        import validate_analysis
        base = _make_base_resume(tmp_path)
        obj = _valid_analysis(base, ['', '', ''])
        errs, _ = validate_analysis.validate(obj)
        assert any('empty' in e for e in errs), f'expected empty-paragraph error, got: {errs}'

    def test_whitespace_paragraphs_rejected(self, tmp_path):
        import validate_analysis
        base = _make_base_resume(tmp_path)
        obj = _valid_analysis(base, ['Real paragraph.', '   ', 'Another real paragraph.'])
        errs, _ = validate_analysis.validate(obj)
        assert any('empty' in e for e in errs), f'expected empty-paragraph error, got: {errs}'

    def test_non_empty_paragraphs_accepted(self, tmp_path):
        import validate_analysis
        base = _make_base_resume(tmp_path)
        obj = _valid_analysis(base, ['First paragraph.', 'Second paragraph.', 'Third paragraph.'])
        errs, _ = validate_analysis.validate(obj)
        assert errs == [], f'unexpected errors: {errs}'


# ── Pipeline failure paths keep analysis.json ─────────────────────────────────

def _make_base_resume(tmp_path) -> Path:
    """Base resume that satisfies the resume quality gate."""
    from docx import Document
    doc = Document()
    doc.add_paragraph('PROFESSIONAL SUMMARY')
    doc.add_paragraph('Original summary line about fund accounting.')
    for section in ('TECHNICAL CAPABILITIES', 'TECHNICAL PROJECTS',
                    'PROFESSIONAL EXPERIENCE', 'EDUCATION'):
        doc.add_paragraph(section)
    for i in range(24):
        doc.add_paragraph(f'Did reconciliation work on fund number {i}.')
    path = tmp_path / 'base_resume.docx'
    doc.save(str(path))
    return path


def _pipeline_env(monkeypatch, tmp_path):
    """Point run_pipeline at a sandbox and neutralize registry + profile IO."""
    run_pipeline = _load_cv_run_pipeline()
    analysis_path = tmp_path / 'analysis.json'
    pending_path = tmp_path / 'pending-analysis.json'
    cv_base = tmp_path / 'CV'
    cv_base.mkdir()
    monkeypatch.setattr(run_pipeline, 'ANALYSIS_PATH', analysis_path)
    monkeypatch.setattr(run_pipeline, 'PENDING_PATH', pending_path)
    monkeypatch.setattr(run_pipeline, 'CV_BASE', cv_base)
    monkeypatch.setattr(run_pipeline, 'rebuild_registry', lambda *a, **k: None)
    monkeypatch.setattr(run_pipeline, '_load_user_profile', lambda: {})
    return run_pipeline, analysis_path


def _valid_analysis(base_path: Path, cover_paragraphs: list[str]) -> dict:
    return {
        'company': 'Acme Fund Services',
        'role': 'Senior Fund Accountant',
        'base_resume_path': str(base_path),
        'summary_edits': [{'old': 'Original summary line about fund accounting.',
                           'new': 'Tailored summary line about fund operations.'}],
        'bullet_edits': [],
        'tailored_edits': [],
        'shared_edits': [],
        'keyword_targets': ['fund accounting'],
        'cover_letter_paragraphs': cover_paragraphs,
        'claims_guardrail': ['Do not invent employers, titles, dates, or certifications.'],
    }


class TestPipelineFailureKeepsAnalysis:
    def test_claims_gate_failure_keeps_analysis_json(self, monkeypatch, tmp_path):
        run_pipeline, analysis_path = _pipeline_env(monkeypatch, tmp_path)
        base = _make_base_resume(tmp_path)
        obj = _valid_analysis(base, ['First paragraph.', 'Second paragraph.', 'Third paragraph.'])
        obj['bullet_edits'] = [{'old': 'Did reconciliation work on fund number 1.',
                                'new': 'Managed $987M AUM single-handedly.'}]
        analysis_path.write_text(json.dumps(obj), encoding='utf-8')

        with pytest.raises(SystemExit) as exc:
            run_pipeline.cmd_apply('Acme Fund Services', 'Senior Fund Accountant')
        assert exc.value.code == 1
        assert analysis_path.exists(), 'analysis.json must survive a claims gate failure'

    def test_empty_cover_letter_body_fails_pipeline(self, monkeypatch, tmp_path):
        """Analysis with 3 empty paragraphs must not produce a body-less letter."""
        run_pipeline, analysis_path = _pipeline_env(monkeypatch, tmp_path)
        base = _make_base_resume(tmp_path)
        obj = _valid_analysis(base, ['', '', ''])
        analysis_path.write_text(json.dumps(obj), encoding='utf-8')

        with pytest.raises(SystemExit) as exc:
            run_pipeline.cmd_apply('Acme Fund Services', 'Senior Fund Accountant')
        assert exc.value.code == 1
        assert analysis_path.exists(), 'analysis.json must survive a validation failure'

    def test_cover_letter_gate_failure_keeps_analysis_json(self, monkeypatch, tmp_path):
        run_pipeline, analysis_path = _pipeline_env(monkeypatch, tmp_path)
        base = _make_base_resume(tmp_path)
        obj = _valid_analysis(base, [
            'I am leveraging my background in fund accounting.',
            'Second paragraph about the team.',
            'Third paragraph, thank you.',
        ])
        analysis_path.write_text(json.dumps(obj), encoding='utf-8')

        manifest = run_pipeline.cmd_apply('Acme Fund Services', 'Senior Fund Accountant')
        assert manifest['status'] == 'fail'
        assert 'cover_letter_qc_failed' in manifest['fail_reasons']
        assert analysis_path.exists(), 'analysis.json must survive a failed run'

    def test_clean_run_passes_and_cleans_up(self, monkeypatch, tmp_path):
        run_pipeline, analysis_path = _pipeline_env(monkeypatch, tmp_path)
        base = _make_base_resume(tmp_path)
        obj = _valid_analysis(base, [
            'I am writing to apply for the role.',
            'My background fits the team well.',
            'Thank you for your consideration.',
        ])
        analysis_path.write_text(json.dumps(obj), encoding='utf-8')

        manifest = run_pipeline.cmd_apply('Acme Fund Services', 'Senior Fund Accountant')
        assert manifest['status'] == 'pass', f"fail_reasons: {manifest['fail_reasons']}"
        assert 'cover_letter_qc' in manifest
        assert manifest['claims_gate']['status'] == 'pass'
        assert not analysis_path.exists(), 'analysis.json should be cleaned up on success'
